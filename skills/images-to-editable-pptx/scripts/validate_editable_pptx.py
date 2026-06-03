from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def validate_pptx(
    pptx_path: Path,
    allow_full_bleed_images: bool = False,
    full_bleed_threshold: float = 0.9,
) -> dict[str, Any]:
    presentation = Presentation(pptx_path)
    slide_area = presentation.slide_width * presentation.slide_height
    failures: list[str] = []
    warnings: list[str] = []
    slides_report: list[dict[str, Any]] = []

    for index, slide in enumerate(presentation.slides, start=1):
        shapes = list(slide.shapes)
        pictures = [shape for shape in shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
        editable = [shape for shape in shapes if shape.shape_type != MSO_SHAPE_TYPE.PICTURE]
        text_shapes = [
            shape
            for shape in editable
            if getattr(shape, "has_text_frame", False)
            and bool(getattr(shape, "text", "").strip())
        ]
        full_bleed_pictures = [
            shape
            for shape in pictures
            if (shape.width * shape.height) / slide_area >= full_bleed_threshold
        ]

        if not shapes:
            failures.append(f"Slide {index}: slide is empty")
        if pictures and not editable:
            failures.append(f"Slide {index}: slide contains pictures only")
        if full_bleed_pictures and not allow_full_bleed_images:
            failures.append(
                f"Slide {index}: contains a full-bleed picture; source slide screenshots are forbidden"
            )
        if len(editable) == 1 and pictures:
            warnings.append(
                f"Slide {index}: only one editable object accompanies picture assets"
            )
        if not text_shapes:
            warnings.append(f"Slide {index}: no editable text was detected")

        slides_report.append(
            {
                "slide": index,
                "shapes": len(shapes),
                "pictures": len(pictures),
                "editable_objects": len(editable),
                "editable_text_objects": len(text_shapes),
                "full_bleed_pictures": len(full_bleed_pictures),
            }
        )

    return {
        "path": str(pptx_path),
        "slide_count": len(presentation.slides),
        "passed": not failures,
        "failures": failures,
        "warnings": warnings,
        "slides": slides_report,
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Reject PPTX decks that flatten slides into full-page images."
    )
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--allow-full-bleed-images", action="store_true")
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()

    report = validate_pptx(args.pptx, args.allow_full_bleed_images)
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"Editable PPTX validation: {'PASS' if report['passed'] else 'FAIL'}")
        print(f"Slides: {report['slide_count']}")
        for failure in report["failures"]:
            print(f"FAIL: {failure}")
        for warning in report["warnings"]:
            print(f"WARN: {warning}")
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
