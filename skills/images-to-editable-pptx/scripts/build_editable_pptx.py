from __future__ import annotations

import argparse
import json
import math
import re
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


SHAPE_TYPES = {
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "chevron": MSO_SHAPE.CHEVRON,
    "diamond": MSO_SHAPE.DIAMOND,
    "hexagon": MSO_SHAPE.HEXAGON,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
}
ALIGNMENTS = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
}
VERTICAL_ALIGNMENTS = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}


def color(value: Any, fallback: str = "#000000") -> RGBColor:
    candidate = str(value or fallback).strip().lstrip("#")
    if not re.fullmatch(r"[0-9a-fA-F]{6}", candidate):
        candidate = fallback.lstrip("#")
    return RGBColor.from_string(candidate.upper())


def inches(value: Any) -> int:
    return Inches(float(value))


def set_shape_name(shape, element: dict[str, Any], fallback: str) -> None:
    shape.name = str(element.get("name") or fallback)


def configure_line(line, element: dict[str, Any], default_color: str = "#000000") -> None:
    line.color.rgb = color(element.get("line") or element.get("color"), default_color)
    line.width = Pt(float(element.get("line_width") or element.get("width") or 1))


def set_dash_and_arrows(shape, element: dict[str, Any]) -> None:
    line_xml = shape._element.spPr.get_or_add_ln()

    dash = str(element.get("dash") or "solid")
    dash_values = {"solid": "solid", "dash": "dash", "dot": "dot"}
    if dash in dash_values and dash != "solid":
        dash_xml = OxmlElement("a:prstDash")
        dash_xml.set("val", dash_values[dash])
        line_xml.append(dash_xml)

    if element.get("arrow_start"):
        head = OxmlElement("a:headEnd")
        head.set("type", "triangle")
        line_xml.append(head)
    if element.get("arrow_end"):
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        line_xml.append(tail)


def configure_fill_and_outline(shape, element: dict[str, Any]) -> None:
    fill_value = element.get("fill")
    if fill_value in {None, "", "none", "transparent"}:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = color(fill_value, "#FFFFFF")

    line_value = element.get("line")
    if line_value in {None, "", "none", "transparent"}:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = color(line_value, "#000000")
        shape.line.width = Pt(float(element.get("line_width") or 1))


def configure_text_frame(
    text_frame,
    element: dict[str, Any],
    theme: dict[str, Any],
) -> None:
    text_frame.clear()
    text_frame.word_wrap = True
    margin = float(element.get("margin") if element.get("margin") is not None else 0.04)
    text_frame.margin_left = inches(margin)
    text_frame.margin_right = inches(margin)
    text_frame.margin_top = inches(margin)
    text_frame.margin_bottom = inches(margin)
    text_frame.vertical_anchor = VERTICAL_ALIGNMENTS.get(
        str(element.get("valign") or "top"), MSO_ANCHOR.TOP
    )

    text = str(element.get("text") or "")
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = ALIGNMENTS.get(
            str(element.get("align") or "left"), PP_ALIGN.LEFT
        )
        paragraph.space_after = Pt(float(element.get("space_after") or 0))
        if not paragraph.runs:
            continue
        run = paragraph.runs[0]
        run.font.name = str(
            element.get("font_face") or theme.get("font_face") or "Microsoft YaHei"
        )
        run.font.size = Pt(float(element.get("font_size") or 18))
        run.font.bold = bool(element.get("bold", False))
        run.font.italic = bool(element.get("italic", False))
        run.font.color.rgb = color(
            element.get("color") or theme.get("text_color"), "#0F172A"
        )


def add_text(slide, element: dict[str, Any], theme: dict[str, Any]) -> None:
    shape = slide.shapes.add_textbox(
        inches(element["x"]),
        inches(element["y"]),
        inches(element["w"]),
        inches(element["h"]),
    )
    set_shape_name(shape, element, "Editable text")
    configure_text_frame(shape.text_frame, element, theme)


def add_shape(slide, element: dict[str, Any], theme: dict[str, Any]) -> None:
    shape_name = str(element.get("shape") or "rectangle")
    if shape_name not in SHAPE_TYPES:
        raise ValueError(f"Unsupported shape: {shape_name}")
    shape = slide.shapes.add_shape(
        SHAPE_TYPES[shape_name],
        inches(element["x"]),
        inches(element["y"]),
        inches(element["w"]),
        inches(element["h"]),
    )
    set_shape_name(shape, element, f"Editable {shape_name}")
    configure_fill_and_outline(shape, element)
    if element.get("text") is not None:
        configure_text_frame(shape.text_frame, element, theme)


def add_line(slide, element: dict[str, Any]) -> None:
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        inches(element["x1"]),
        inches(element["y1"]),
        inches(element["x2"]),
        inches(element["y2"]),
    )
    set_shape_name(shape, element, "Editable line")
    configure_line(shape.line, element)
    set_dash_and_arrows(shape, element)


def add_image(
    slide,
    element: dict[str, Any],
    base_dir: Path,
    source_images: set[Path],
    slide_area: float,
    allow_full_bleed_images: bool,
) -> None:
    image_path = (base_dir / str(element["path"])).resolve()
    if image_path in source_images:
        raise ValueError(f"Refusing to embed source slide reference image: {image_path}")
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    area = float(element["w"]) * float(element["h"])
    if area / slide_area >= 0.9 and not allow_full_bleed_images:
        raise ValueError(
            f"Refusing full-bleed image asset without explicit approval: {image_path}"
        )
    shape = slide.shapes.add_picture(
        str(image_path),
        inches(element["x"]),
        inches(element["y"]),
        width=inches(element["w"]),
        height=inches(element["h"]),
    )
    set_shape_name(shape, element, "Bounded image asset")


def set_cell_text(cell, text: str, element: dict[str, Any], theme: dict[str, Any], header: bool) -> None:
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = color(
        element.get("header_fill" if header else "body_fill"),
        "#0D9488" if header else "#FFFFFF",
    )
    text_frame = cell.text_frame
    text_frame.margin_left = inches(0.06)
    text_frame.margin_right = inches(0.06)
    text_frame.margin_top = inches(0.04)
    text_frame.margin_bottom = inches(0.04)
    for paragraph in text_frame.paragraphs:
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.name = str(theme.get("font_face") or "Microsoft YaHei")
            run.font.size = Pt(float(element.get("font_size") or 13))
            run.font.bold = header
            run.font.color.rgb = color(
                element.get("header_color" if header else "body_color"),
                "#FFFFFF" if header else "#0F172A",
            )


def add_table(slide, element: dict[str, Any], theme: dict[str, Any]) -> None:
    rows = element.get("rows")
    if not isinstance(rows, list) or not rows or not isinstance(rows[0], list):
        raise ValueError("Table rows must be a non-empty list of lists")
    column_count = len(rows[0])
    if column_count == 0 or any(len(row) != column_count for row in rows):
        raise ValueError("Every table row must have the same number of columns")

    frame = slide.shapes.add_table(
        len(rows),
        column_count,
        inches(element["x"]),
        inches(element["y"]),
        inches(element["w"]),
        inches(element["h"]),
    )
    set_shape_name(frame, element, "Editable table")
    table = frame.table
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            set_cell_text(
                table.cell(row_index, column_index),
                str(value),
                element,
                theme,
                header=row_index == 0,
            )


def add_small_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    size: float,
    value_color: Any,
    theme: dict[str, Any],
    align: str = "center",
    bold: bool = False,
    name: str = "Editable chart label",
) -> None:
    add_text(
        slide,
        {
            "type": "text",
            "name": name,
            "x": x,
            "y": y,
            "w": w,
            "h": h,
            "text": text,
            "font_size": size,
            "color": value_color,
            "align": align,
            "valign": "middle",
            "bold": bold,
            "margin": 0,
        },
        theme,
    )


def add_bar_chart(slide, element: dict[str, Any], theme: dict[str, Any]) -> None:
    data = element.get("data")
    if not isinstance(data, list) or not data:
        raise ValueError("bar_chart data must be a non-empty list")
    values = [float(item["value"]) for item in data]
    max_value = max(values)
    if max_value <= 0:
        raise ValueError("bar_chart values must include a positive value")

    x, y, w, h = map(float, (element["x"], element["y"], element["w"], element["h"]))
    label_h = 0.35
    value_h = 0.35
    chart_top = y + value_h
    chart_bottom = y + h - label_h
    chart_height = chart_bottom - chart_top
    slot = w / len(data)
    bar_w = min(slot * 0.55, 0.8)
    suffix = str(element.get("value_suffix") or "")
    label_color = element.get("label_color") or "#475569"
    value_color = element.get("value_color") or theme.get("text_color") or "#0F172A"

    add_line(
        slide,
        {
            "name": "Editable bar chart baseline",
            "x1": x,
            "y1": chart_bottom,
            "x2": x + w,
            "y2": chart_bottom,
            "color": element.get("axis_color") or "#CBD5E1",
            "width": 1,
        },
    )
    for index, item in enumerate(data):
        value = float(item["value"])
        bar_height = max(0.04, chart_height * value / max_value)
        bar_x = x + index * slot + (slot - bar_w) / 2
        bar_y = chart_bottom - bar_height
        add_shape(
            slide,
            {
                "type": "shape",
                "shape": "rectangle",
                "name": f"Editable bar {index + 1}",
                "x": bar_x,
                "y": bar_y,
                "w": bar_w,
                "h": bar_height,
                "fill": item.get("color") or element.get("bar_color") or "#0D9488",
                "line": "none",
            },
            theme,
        )
        add_small_text(
            slide,
            f"{item['value']}{suffix}",
            bar_x - 0.15,
            bar_y - 0.34,
            bar_w + 0.3,
            0.28,
            11,
            value_color,
            theme,
            bold=True,
            name=f"Editable bar value {index + 1}",
        )
        add_small_text(
            slide,
            str(item.get("label") or ""),
            x + index * slot,
            chart_bottom + 0.03,
            slot,
            0.28,
            10,
            label_color,
            theme,
            name=f"Editable bar label {index + 1}",
        )


def add_line_chart(slide, element: dict[str, Any], theme: dict[str, Any]) -> None:
    data = element.get("data")
    if not isinstance(data, list) or len(data) < 2:
        raise ValueError("line_chart data must contain at least two points")
    values = [float(item["value"]) for item in data]
    min_value, max_value = min(values), max(values)
    value_range = max(max_value - min_value, 1)
    x, y, w, h = map(float, (element["x"], element["y"], element["w"], element["h"]))
    label_h = 0.35
    value_h = 0.35
    chart_top = y + value_h
    chart_bottom = y + h - label_h
    chart_height = chart_bottom - chart_top
    step = w / (len(data) - 1)
    suffix = str(element.get("value_suffix") or "")
    line_color = element.get("line_color") or "#0D9488"
    marker_color = element.get("marker_color") or line_color
    label_color = element.get("label_color") or "#475569"
    value_color = element.get("value_color") or theme.get("text_color") or "#0F172A"

    points: list[tuple[float, float]] = []
    for index, value in enumerate(values):
        point_x = x + index * step
        point_y = chart_bottom - (value - min_value) / value_range * chart_height
        points.append((point_x, point_y))

    add_line(
        slide,
        {
            "name": "Editable line chart baseline",
            "x1": x,
            "y1": chart_bottom,
            "x2": x + w,
            "y2": chart_bottom,
            "color": element.get("axis_color") or "#CBD5E1",
            "width": 1,
        },
    )
    for index in range(len(points) - 1):
        start, end = points[index], points[index + 1]
        add_line(
            slide,
            {
                "name": f"Editable trend segment {index + 1}",
                "x1": start[0],
                "y1": start[1],
                "x2": end[0],
                "y2": end[1],
                "color": line_color,
                "width": float(element.get("line_width") or 2.5),
            },
        )
    marker_size = float(element.get("marker_size") or 0.13)
    for index, (item, point) in enumerate(zip(data, points, strict=True)):
        add_shape(
            slide,
            {
                "type": "shape",
                "shape": "ellipse",
                "name": f"Editable trend marker {index + 1}",
                "x": point[0] - marker_size / 2,
                "y": point[1] - marker_size / 2,
                "w": marker_size,
                "h": marker_size,
                "fill": marker_color,
                "line": marker_color,
            },
            theme,
        )
        add_small_text(
            slide,
            f"{item['value']}{suffix}",
            point[0] - 0.4,
            point[1] - 0.42,
            0.8,
            0.25,
            10,
            value_color,
            theme,
            bold=True,
            name=f"Editable trend value {index + 1}",
        )
        add_small_text(
            slide,
            str(item.get("label") or ""),
            point[0] - 0.4,
            chart_bottom + 0.03,
            0.8,
            0.25,
            10,
            label_color,
            theme,
            name=f"Editable trend label {index + 1}",
        )


def resolve_source_images(spec: dict[str, Any], base_dir: Path) -> set[Path]:
    paths: set[Path] = set()
    for slide in spec.get("slides") or []:
        source_image = slide.get("source_image")
        if source_image:
            paths.add((base_dir / str(source_image)).resolve())
    return paths


def build_pptx(
    spec: dict[str, Any],
    output_path: Path,
    base_dir: Path,
    allow_full_bleed_images: bool = False,
) -> Path:
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("Spec must contain at least one slide")

    slide_size = spec.get("slide_size") or {}
    slide_width = float(slide_size.get("width") or 13.333)
    slide_height = float(slide_size.get("height") or 7.5)
    slide_area = slide_width * slide_height
    theme = spec.get("theme") if isinstance(spec.get("theme"), dict) else {}
    source_images = resolve_source_images(spec, base_dir)

    presentation = Presentation()
    presentation.slide_width = inches(slide_width)
    presentation.slide_height = inches(slide_height)
    presentation.core_properties.title = str(spec.get("title") or output_path.stem)
    blank_layout = presentation.slide_layouts[6]

    for slide_index, slide_spec in enumerate(slides, start=1):
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = color(
            slide_spec.get("background"), "#FFFFFF"
        )
        elements = slide_spec.get("elements")
        if not isinstance(elements, list):
            raise ValueError(f"Slide {slide_index} elements must be a list")

        for element_index, element in enumerate(elements, start=1):
            if not isinstance(element, dict):
                raise ValueError(
                    f"Slide {slide_index} element {element_index} must be an object"
                )
            element_type = str(element.get("type") or "")
            try:
                if element_type == "text":
                    add_text(slide, element, theme)
                elif element_type == "shape":
                    add_shape(slide, element, theme)
                elif element_type == "line":
                    add_line(slide, element)
                elif element_type == "image":
                    add_image(
                        slide,
                        element,
                        base_dir,
                        source_images,
                        slide_area,
                        allow_full_bleed_images,
                    )
                elif element_type == "table":
                    add_table(slide, element, theme)
                elif element_type == "bar_chart":
                    add_bar_chart(slide, element, theme)
                elif element_type == "line_chart":
                    add_line_chart(slide, element, theme)
                else:
                    raise ValueError(f"Unsupported element type: {element_type}")
            except Exception as exc:
                raise ValueError(
                    f"Slide {slide_index} element {element_index} failed: {exc}"
                ) from exc

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Build a genuinely editable PPTX from a native reconstruction spec."
    )
    parser.add_argument("--spec", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--allow-full-bleed-images", action="store_true")
    args = parser.parse_args()

    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    build_pptx(
        spec=spec,
        output_path=args.out,
        base_dir=args.spec.parent,
        allow_full_bleed_images=args.allow_full_bleed_images,
    )
    print(f"Created editable PPTX: {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
