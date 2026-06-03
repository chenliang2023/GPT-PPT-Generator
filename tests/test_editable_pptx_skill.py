from __future__ import annotations

import importlib.util
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches


ROOT = Path(__file__).resolve().parents[1]
SKILL_SCRIPTS = ROOT / "skills" / "images-to-editable-pptx" / "scripts"


def load_script(name: str):
    path = SKILL_SCRIPTS / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_build_and_validate_editable_deck(tmp_path: Path) -> None:
    builder = load_script("build_editable_pptx")
    validator = load_script("validate_editable_pptx")
    output = tmp_path / "editable.pptx"
    spec = {
        "title": "Editable Test",
        "slides": [
            {
                "background": "#F8FAFC",
                "elements": [
                    {
                        "type": "text",
                        "name": "Title",
                        "x": 0.8,
                        "y": 0.6,
                        "w": 6.0,
                        "h": 0.8,
                        "text": "Editable title",
                        "font_size": 30,
                        "bold": True,
                    },
                    {
                        "type": "bar_chart",
                        "x": 0.8,
                        "y": 2.0,
                        "w": 6.5,
                        "h": 4.2,
                        "data": [
                            {"label": "A", "value": 42},
                            {"label": "B", "value": 68},
                        ],
                    },
                ],
            }
        ],
    }

    builder.build_pptx(spec, output, tmp_path)
    report = validator.validate_pptx(output)
    presentation = Presentation(output)

    assert report["passed"] is True
    assert len(presentation.slides) == 1
    assert all(
        shape.shape_type != MSO_SHAPE_TYPE.PICTURE
        for shape in presentation.slides[0].shapes
    )


def test_validator_rejects_full_page_picture(tmp_path: Path) -> None:
    validator = load_script("validate_editable_pptx")
    image_path = tmp_path / "slide.png"
    Image.new("RGB", (1600, 900), (255, 255, 255)).save(image_path)
    output = tmp_path / "flattened.pptx"

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(image_path),
        0,
        0,
        width=presentation.slide_width,
        height=presentation.slide_height,
    )
    presentation.save(output)

    report = validator.validate_pptx(output)
    assert report["passed"] is False
    assert any("pictures only" in failure for failure in report["failures"])
